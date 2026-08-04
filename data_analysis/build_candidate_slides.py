"""Build a PPTX deck for Figure 4C Jurkat candidate review.

One slide per candidate with:
  - Title (gene + UniProt)
  - Pattern one-liner
  - Candidate plot on the left (PNG from Figure4C_candidate_slides.R)
  - Pros (green) and Cons (red) bullets on the right
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
FIG_DIR = Path(
    "/Volumes/cos-lab-rwu60/Longping/OGlycoTM_Final_Version/Figures/Figure4/Figure4C_candidate_slides"
)
OUT_PPTX = FIG_DIR / "Figure4C_candidate_review.pptx"

# ---------------------------------------------------------------------------
# Candidate data (order = how Ronghu will see them)
# ---------------------------------------------------------------------------
CANDIDATES = [
    {
        "gene": "CCAR1",
        "uniprot": "Q8IX12",
        "pattern": "HEK -0.95, HepG2 -0.25, Jurkat +0.68  (3 facets)",
        "pros": [
            "3 facets; best fit to current Figure 4C concept",
            "Localized sites in all three cell types",
            "Jurkat site-DE exists",
            "One exact site (S287) is in OGlcNAc Atlas",
            "Already used in current Figure 4C script",
        ],
        "cons": [
            "HepG2 reproducibility is the weak point: replicate log2FCs "
            "are -0.46, +0.07, -0.41",
            "Biology is only indirectly connected to Figure 4B, not a clean "
            "immune marker",
        ],
    },
    {
        "gene": "YIPF3",
        "uniprot": "Q9GZM5",
        "pattern": "HepG2 -0.51, Jurkat +0.90  (2 facets)",
        "pros": [
            "Cleanest data of the non-Jurkat-only options",
            "Very tight replicates",
            "Localized S349 in both HepG2 and Jurkat",
            "Site-DE in both HepG2 and Jurkat",
            "Exact S349 already reported in OGlcNAc Atlas",
        ],
        "cons": [
            "Only 2 facets (not quantified in HEK293T)",
            "Pushes the panel further toward vesicle/secretory trafficking, "
            "which overlaps with SEC31A",
        ],
    },
    {
        "gene": "CIC",
        "uniprot": "Q96RK0",
        "pattern": "HEK -0.24, HepG2 +0.06, Jurkat +0.66  (3 facets)",
        "pros": [
            "Best visual replacement for the old CTTN pattern",
            "3 facets",
            "Very tight replicates in all three cells",
        ],
        "cons": [
            "No localized Jurkat site in dataset",
            "No Jurkat site-DE",
            "Weak connection to the Figure 4B Jurkat story",
        ],
    },
    {
        "gene": "YIF1B",
        "uniprot": "Q5BJH7",
        "pattern": "HEK +0.67, HepG2 +0.02, Jurkat +0.76  (3 facets)",
        "pros": [
            "3 facets",
            "Clean plot",
            "Jurkat has localized S28",
            "Exact S28 is in OGlcNAc Atlas",
        ],
        "cons": [
            "Not Jurkat-specific (HEK is also clearly up)",
            "Jurkat site-DE is not significant",
            "Another trafficking-flavored protein",
        ],
    },
]

# ---------------------------------------------------------------------------
# Colors
# ---------------------------------------------------------------------------
GREEN = RGBColor(0x1F, 0x7A, 0x32)   # pros
RED   = RGBColor(0xB0, 0x28, 0x28)   # cons
DARK  = RGBColor(0x22, 0x22, 0x22)   # body text
GREY  = RGBColor(0x66, 0x66, 0x66)   # pattern subtitle


def _set_font(run, *, size, bold=False, color=DARK, name="Arial"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_candidate_slide(prs: Presentation, cand: dict) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # ---- Title ----
    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(12.6), Inches(0.55))
    tf = title_box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"{cand['gene']}  ({cand['uniprot']})"
    _set_font(r, size=28, bold=True, color=DARK)

    # ---- Subtitle (pattern) ----
    sub_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.80), Inches(12.6), Inches(0.4))
    sf = sub_box.text_frame
    sf.margin_left = sf.margin_right = sf.margin_top = sf.margin_bottom = 0
    sp = sf.paragraphs[0]
    sp.alignment = PP_ALIGN.LEFT
    sr = sp.add_run()
    sr.text = cand["pattern"]
    _set_font(sr, size=16, bold=False, color=GREY)

    # ---- Figure (left) ----
    png_path = FIG_DIR / f"{cand['gene']}_{cand['uniprot']}.png"
    if not png_path.exists():
        raise FileNotFoundError(png_path)

    # Image area: left half of slide
    img_left = Inches(0.4)
    img_top = Inches(1.4)
    img_max_h = Inches(5.6)
    img_max_w = Inches(5.8)
    pic = slide.shapes.add_picture(str(png_path), img_left, img_top, height=img_max_h)
    # Scale down if it exceeds max width
    if pic.width > img_max_w:
        scale = img_max_w / pic.width
        pic.width = int(pic.width * scale)
        pic.height = int(pic.height * scale)
    # Vertically center image within allotted space
    cy = Inches(1.4) + (img_max_h - pic.height) // 2
    pic.top = cy
    # Horizontally center image within left column (0.4 to 6.3)
    col_center = (Inches(0.4) + Inches(6.3)) // 2
    pic.left = col_center - pic.width // 2

    # ---- Pros / Cons (right) ----
    text_left = Inches(6.6)
    text_top = Inches(1.4)
    text_w = Inches(6.5)
    text_h = Inches(5.6)
    tbox = slide.shapes.add_textbox(text_left, text_top, text_w, text_h)
    ttf = tbox.text_frame
    ttf.word_wrap = True
    ttf.margin_left = ttf.margin_right = Inches(0.05)
    ttf.margin_top = ttf.margin_bottom = Inches(0.05)

    # Pros header
    p = ttf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = "Pros"
    _set_font(r, size=20, bold=True, color=GREEN)

    for item in cand["pros"]:
        para = ttf.add_paragraph()
        para.level = 0
        para.space_before = Pt(2)
        para.space_after = Pt(2)
        rr = para.add_run()
        rr.text = f"•  {item}"
        _set_font(rr, size=14, bold=False, color=DARK)

    # Spacer
    spacer = ttf.add_paragraph()
    spacer.space_before = Pt(8)
    spacer.space_after = Pt(2)
    sr = spacer.add_run()
    sr.text = ""
    _set_font(sr, size=6, color=DARK)

    # Cons header
    cons_head = ttf.add_paragraph()
    cons_head.space_before = Pt(6)
    cons_head.space_after = Pt(2)
    cr = cons_head.add_run()
    cr.text = "Cons"
    _set_font(cr, size=20, bold=True, color=RED)

    for item in cand["cons"]:
        para = ttf.add_paragraph()
        para.level = 0
        para.space_before = Pt(2)
        para.space_after = Pt(2)
        rr = para.add_run()
        rr.text = f"•  {item}"
        _set_font(rr, size=14, bold=False, color=DARK)


def add_title_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(2.6), Inches(12.3), Inches(1.2))
    tf = tbox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Figure 4C Jurkat Example Candidates"
    _set_font(r, size=36, bold=True, color=DARK)

    sbox = slide.shapes.add_textbox(Inches(0.5), Inches(3.8), Inches(12.3), Inches(0.8))
    sf = sbox.text_frame
    sf.word_wrap = True
    sp = sf.paragraphs[0]
    sp.alignment = PP_ALIGN.CENTER
    sr = sp.add_run()
    sr.text = "Please pick a replacement for RAC2. Each slide = 1 candidate."
    _set_font(sr, size=20, bold=False, color=GREY)


def main() -> None:
    prs = Presentation()
    # 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)
    for cand in CANDIDATES:
        add_candidate_slide(prs, cand)

    prs.save(str(OUT_PPTX))
    print(f"Saved: {OUT_PPTX}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    main()
